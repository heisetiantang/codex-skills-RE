# -*- coding: utf-8 -*-
"""餐食记录追加脚本：把一次餐食分析结果追加到 Excel + Word，并处理新增食物。
用法:
  python append_meal.py <meal.json>
meal.json 结构:
{
  "date": "2026-08-11", "meal": "早餐", "detail": "食物明细文字",
  "kcal": 370, "carb": 35, "protein": 25, "fat": 14, "fiber": 5,
  "va": 200, "vb1": 0.15, "vb2": 0.30, "vc": 35, "ve": 3,
  "note": "备注", "advice": ["建议1", "建议2"],
  "new_foods": [["测试饼", "100g", 300, 55, 6, 5, 2, 0, 0.05, 0.05, 0, 0.5, "TEST"]]
}
- new_foods（可选）: 本餐新识别、需追加到食物数据库的食物行（列同数据库表头）；
  同名同单位已存在则跳过；成功新增的会写进备注标记【新增食物:...】，撤回时可回滚。
photo（可选）: 照片原始路径，将复制到 health/photos/日期_餐次.jpg
"""
import sys, os, json, shutil
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from docx import Document
from docx.shared import Pt

BASE = os.path.dirname(os.path.abspath(__file__))
from _config import DATA_DIR, PROFILE
REC_XLSX = os.path.join(DATA_DIR, "餐食营养记录.xlsx")
REC_DOCX = os.path.join(DATA_DIR, "餐食营养记录.docx")
DB_XLSX = os.path.join(DATA_DIR, "食物营养数据库.xlsx")
PHOTO_DIR = os.path.join(DATA_DIR, "photos")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="C9C9C9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def append_new_foods(m):
    """把本餐新增食物写入数据库，返回实际新增的食物名列表（同名同单位跳过）。"""
    new_foods = m.get("new_foods") or []
    if not new_foods:
        return []
    wb = load_workbook(DB_XLSX)
    ws = wb["食物数据库"]
    existing = {(str(r[0]), str(r[1])) for r in ws.iter_rows(min_row=2, max_col=2, values_only=True)}
    added = []
    for row in new_foods:
        key = (str(row[0]), str(row[1]))
        if key in existing:
            print("SKIP existing food:", row[0], row[1])
            continue
        ws.append(row)
        existing.add(key)
        added.append(str(row[0]))
    wb.save(DB_XLSX)
    if added:
        print("DB foods added:", added)
    return added


def append_xlsx(m, added_foods):
    wb = load_workbook(REC_XLSX)
    ws = wb["餐食记录"]
    note = m.get("note", "")
    if added_foods:
        note = (note + " " if note else "") + "【新增食物:" + ",".join(added_foods) + "】"
    row = [m["date"], m["meal"], m["detail"], m["kcal"], m["carb"], m["protein"],
           m["fat"], m["fiber"], m.get("va", "-"), m.get("vb1", "-"),
           m.get("vb2", "-"), m.get("vc", "-"), m.get("ve", "-"), "±20%", note]
    ws.append(row)
    r = ws.max_row
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row=r, column=c)
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = LEFT if c == 3 else CENTER
    wb.save(REC_XLSX)
    print("XLSX row appended at row", r)


def append_docx(m):
    doc = Document(REC_DOCX)
    doc.add_heading(f"{m['date']} {m['meal']}", level=1)
    doc.add_paragraph(f"食物明细：{m['detail']}")
    doc.add_heading("营养汇总", level=2)
    t = doc.add_table(rows=1, cols=2)
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    hdr[0].text, hdr[1].text = "指标", "数值"
    rows = [
        ("总热量", f"{m['kcal']} kcal"),
        ("碳水 / 供能比", f"{m['carb']} g / {round(m['carb'] * 4 / max(m['kcal'], 1) * 100)}%"),
        ("蛋白质 / 供能比", f"{m['protein']} g / {round(m['protein'] * 4 / max(m['kcal'], 1) * 100)}%"),
        ("脂肪 / 供能比", f"{m['fat']} g / {round(m['fat'] * 9 / max(m['kcal'], 1) * 100)}%"),
        ("膳食纤维", f"{m['fiber']} g"),
        ("维生素A", f"{m.get('va', '-')} ug"),
        ("维生素B1 / B2", f"{m.get('vb1', '-')} / {m.get('vb2', '-')} mg"),
        ("维生素C / E", f"{m.get('vc', '-')} / {m.get('ve', '-')} mg"),
    ]
    for k, v in rows:
        c = t.add_row().cells
        c[0].text, c[1].text = k, v
    advice = m.get("advice") or []
    if advice:
        doc.add_heading("建议", level=2)
        for s in advice:
            doc.add_paragraph(s, style="List Bullet")
    doc.save(REC_DOCX)
    print("DOCX section appended")


def copy_photo(m):
    photo = m.get("photo")
    if not photo or not os.path.exists(photo):
        return
    os.makedirs(PHOTO_DIR, exist_ok=True)
    ext = os.path.splitext(photo)[1] or ".jpg"
    base = os.path.join(PHOTO_DIR, f"{m['date']}_{m['meal']}")
    dest = base + ext
    n = 2
    while os.path.exists(dest):  # 同日同餐次多张照片自动加序号，不覆盖旧照片
        dest = f"{base}_{n}{ext}"
        n += 1
    shutil.copy2(photo, dest)
    print("PHOTO saved:", dest)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("usage: python append_meal.py <meal.json>")
        sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as f:
        meal = json.load(f)
    added = append_new_foods(meal)
    append_xlsx(meal, added)
    append_docx(meal)
    copy_photo(meal)
    print("DONE")
