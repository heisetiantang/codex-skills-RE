# -*- coding: utf-8 -*-
"""初始化餐食营养记录系统：创建三个核心文件并写入首条记录。
- 餐食营养记录.xlsx   （Sheet: 餐食记录 + 说明）
- 餐食营养记录.docx   （可读归档）
- 食物营养数据库.xlsx  （Sheet: 食物数据库 + 说明，约34种常用食物）
数据参照《中国食物成分表》标准版，均为参考值。
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = os.path.dirname(os.path.abspath(__file__))
from _config import DATA_DIR, PROFILE
REC_XLSX = os.path.join(DATA_DIR, "餐食营养记录.xlsx")
DB_XLSX = os.path.join(DATA_DIR, "食物营养数据库.xlsx")
REC_DOCX = os.path.join(DATA_DIR, "餐食营养记录.docx")

HDR_FILL = PatternFill("solid", fgColor="2F5B66")
HDR_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="C9C9C9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

# 食物数据库列
DB_HEADERS = ["食物", "单位", "热量(kcal)", "碳水(g)", "蛋白(g)", "脂肪(g)",
              "纤维(g)", "VA(ug)", "VB1(mg)", "VB2(mg)", "VC(mg)", "VE(mg)", "备注"]
# 单位: 每 100g / 100ml / 1个 / 1根 / 1片 / 10g 等
FOODS = [
    ["鸡蛋(白煮)", "100g", 144, 2.8, 12.7, 9.0, 0, 234, 0.11, 0.32, 0, 1.84, "可食部约88%"],
    ["鸡蛋(白煮)", "1个(50g)", 72, 1.4, 6.4, 4.5, 0, 117, 0.06, 0.16, 0, 0.90, "带壳熟蛋约50-55g"],
    ["鸡蛋清", "1个(35g)", 21, 1.1, 4.1, 0.0, 0, 0, 0.01, 0.11, 0, 0.00, "几乎无脂肪"],
    ["鸡蛋黄", "1个(15g)", 49, 0.5, 2.3, 4.2, 0, 66, 0.05, 0.04, 0, 0.76, "胆固醇约180mg"],
    ["茶叶蛋(卤蛋)", "1个(55g)", 77, 1.5, 6.6, 5.0, 0, 129, 0.06, 0.18, 0, 1.00, "按全蛋近似"],
    ["煎蛋(油煎)", "1个(45g)", 105, 0.9, 6.0, 8.5, 0, 180, 0.05, 0.14, 0, 1.50, "含煎油"],
    ["无糖豆浆", "100ml", 14, 0.5, 1.8, 0.7, 0.2, 0, 0.02, 0.02, 0, 0.80, "不加糖"],
    ["玉米面窝窝头", "1个(80g)", 174, 32.6, 6.1, 3.1, 2.6, 16, 0.10, 0.07, 0, 0.50, "100g≈217kcal"],
    ["小米粥", "100g", 46, 8.4, 1.4, 0.7, 0.9, 3, 0.02, 0.03, 0, 0.10, "稠粥"],
    ["白米粥", "100g", 46, 9.9, 1.1, 0.3, 0.3, 0, 0.01, 0.02, 0, 0.10, ""],
    ["馒头", "100g", 223, 47.0, 7.0, 1.1, 1.3, 0, 0.04, 0.05, 0, 0.30, ""],
    ["花卷", "100g", 217, 45.6, 6.4, 1.0, 1.5, 0, 0.03, 0.05, 0, 0.30, ""],
    ["猪肉包子", "1个(80g)", 200, 26.0, 7.0, 8.0, 1.2, 10, 0.08, 0.07, 1, 0.50, "估算"],
    ["油条", "100g", 388, 46.0, 6.9, 17.6, 0.9, 0, 0.05, 0.06, 0, 3.20, "1根约60g"],
    ["咸鸭蛋", "1个(60g)", 114, 3.4, 8.0, 9.0, 0, 174, 0.09, 0.26, 0, 3.00, "高钠"],
    ["甜玉米(煮)", "100g可食", 112, 22.8, 4.0, 1.2, 2.9, 8, 0.03, 0.06, 8, 0.10, "1根可食约150g"],
    ["红薯(煮)", "100g", 106, 25.2, 1.1, 0.2, 1.6, 125, 0.04, 0.04, 26, 0.30, "1个约150g"],
    ["全麦面包", "1片(35g)", 85, 15.0, 3.5, 1.2, 2.2, 0, 0.05, 0.03, 0, 0.30, ""],
    ["燕麦片(干)", "100g", 367, 61.6, 15.0, 6.7, 5.3, 0, 0.30, 0.13, 0, 0.80, "冲调另计"],
    ["牛奶(全脂)", "100ml", 65, 4.9, 3.3, 3.6, 0, 24, 0.03, 0.14, 1, 0.10, ""],
    ["酸奶(原味)", "100g", 72, 9.3, 2.5, 2.7, 0, 26, 0.03, 0.15, 1, 0.10, "含糖型更高"],
    ["豆腐脑(咸)", "100g", 47, 3.0, 5.0, 1.9, 0.3, 10, 0.03, 0.05, 0, 1.00, "估算"],
    ["豆腐(北)", "100g", 98, 3.9, 9.2, 5.9, 0.5, 10, 0.05, 0.03, 0, 3.10, ""],
    ["凉拌大白菜", "100g", 17, 3.2, 1.5, 0.1, 0.8, 20, 0.04, 0.05, 31, 0.40, "生重,含拌油另计"],
    ["水发木耳", "100g", 27, 6.0, 1.5, 0.2, 2.6, 3, 0.02, 0.05, 0, 0.40, ""],
    ["枸杞(干)", "10g", 26, 6.4, 1.4, 0.2, 1.6, 162, 0.04, 0.05, 1, 0.20, "100g VA≈1620ug"],
    ["凉拌黄瓜", "100g", 15, 2.9, 0.8, 0.2, 0.5, 15, 0.02, 0.03, 9, 0.50, ""],
    ["小番茄", "100g", 25, 5.8, 0.9, 0.2, 1.2, 42, 0.03, 0.03, 19, 0.60, ""],
    ["香蕉", "1根(120g可食)", 108, 26.0, 1.4, 0.2, 1.5, 8, 0.02, 0.04, 10, 0.20, ""],
    ["苹果", "1个(200g可食)", 106, 27.0, 0.4, 0.4, 4.8, 8, 0.03, 0.04, 8, 0.40, ""],
    ["花生米(熟)", "10g", 60, 2.2, 2.4, 4.4, 0.7, 1, 0.03, 0.02, 0, 1.00, "100g≈600kcal"],
    ["核桃(干)", "10g", 65, 1.0, 1.5, 6.0, 0.9, 1, 0.02, 0.02, 0, 0.50, "100g≈646kcal"],
    ["酱菜/咸菜", "10g", 4, 0.8, 0.2, 0.1, 0.3, 0, 0.01, 0.01, 0, 0.00, "高钠,适量"],
    ["调味油(凉拌)", "5g", 45, 0.0, 0.0, 5.0, 0, 0, 0.00, 0.00, 0, 0.50, "植物油"],
]

# 首条餐食记录（2026-08-10 早餐 · 黄金模板）
FIRST_MEAL = {
    "date": "2026-08-10",
    "meal": "早餐",
    "detail": "凉拌菜(大白菜+水发木耳+枸杞)约150g；茶叶蛋1个(仅蛋白约30g)；白煮蛋1个(1蛋黄+1蛋白约50g)；玉米面窝窝头1个80g；无糖豆浆200ml",
    "kcal": 370, "carb": 35, "protein": 25, "fat": 14, "fiber": 5,
    "va": 200, "vb1": 0.15, "vb2": 0.30, "vc": 35, "ve": 3,
    "note": "黄金模板；实际吃法=1蛋黄+2蛋白；维生素为估算",
}


def style_sheet(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = LEFT if cell.column == 1 else CENTER
    ws.freeze_panes = "A2"


def build_db():
    wb = Workbook()
    ws = wb.active
    ws.title = "食物数据库"
    ws.append(DB_HEADERS)
    for f in FOODS:
        ws.append(f)
    style_sheet(ws, [16, 13, 10, 9, 9, 9, 9, 9, 9, 9, 9, 9, 26])
    ws2 = wb.create_sheet("说明")
    notes = [
        ["说明项", "内容"],
        ["数据来源", "《中国食物成分表》标准版参考值"],
        ["单位口径", "每100g/100ml/1个/1根/1片/10g 可食部"],
        ["误差", "克重与营养素均为估算,误差约±20%"],
        ["更新方式", "遇到新食物时自动追加行,不覆盖旧数据"],
        ["维生素口径", "VA为ug;VB1/VB2/VC/VE为mg;查不到标\"-\""],
    ]
    for n in notes:
        ws2.append(n)
    style_sheet(ws2, [14, 60])
    wb.save(DB_XLSX)
    print("DB OK:", DB_XLSX)


def build_records_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "餐食记录"
    headers = ["日期", "餐次", "食物明细", "总热量(kcal)", "碳水(g)", "蛋白(g)", "脂肪(g)",
               "纤维(g)", "VA(ug)", "VB1(mg)", "VB2(mg)", "VC(mg)", "VE(mg)", "误差", "备注"]
    ws.append(headers)
    m = FIRST_MEAL
    ws.append([m["date"], m["meal"], m["detail"], m["kcal"], m["carb"], m["protein"],
               m["fat"], m["fiber"], m["va"], m["vb1"], m["vb2"], m["vc"], m["ve"],
               "±20%", m["note"]])
    style_sheet(ws, [12, 8, 52, 12, 9, 9, 9, 9, 9, 9, 9, 9, 9, 9, 34])
    ws2 = wb.create_sheet("说明")
    notes = [
        ["说明项", "内容"],
        ["记录方式", "每次发餐食照片追加一行,不覆盖历史"],
        ["餐次判定", "6:00-9:30早餐 / 11:00-13:30午餐 / 17:00-19:30晚餐 / 其余加餐;用户明说加餐优先"],
        ["尺寸参照", "食堂餐盘短边20cm、碗直径10cm"],
        ["误差", "克重与营养素为估算,约±20%"],
        ["维生素口径", "VA为ug;VB1/VB2/VC/VE为mg;查不到标\"-\""],
    ]
    for n in notes:
        ws2.append(n)
    style_sheet(ws2, [14, 60])
    wb.save(REC_XLSX)
    print("REC XLSX OK:", REC_XLSX)


def build_records_docx():
    doc = Document()
    doc.styles["Normal"].font.name = "微软雅黑"
    doc.styles["Normal"].font.size = Pt(10.5)
    h = doc.add_heading("餐食营养记录", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("记录方式：每次餐食照片追加一节 · 克重与营养素为估算（误差约±20%）")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    m = FIRST_MEAL
    doc.add_heading(f"{m['date']} {m['meal']}", level=1)
    doc.add_paragraph(f"食物明细：{m['detail']}")

    doc.add_heading("营养汇总", level=2)
    t = doc.add_table(rows=1, cols=2)
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    hdr[0].text, hdr[1].text = "指标", "数值"
    rows = [
        ("总热量", f"{m['kcal']} kcal"),
        ("碳水 / 供能比", f"{m['carb']} g / 38%"),
        ("蛋白质 / 供能比", f"{m['protein']} g / 28%"),
        ("脂肪 / 供能比", f"{m['fat']} g / 34%"),
        ("膳食纤维", f"{m['fiber']} g"),
        ("维生素A", f"{m['va']} ug"),
        ("维生素B1 / B2", f"{m['vb1']} / {m['vb2']} mg"),
        ("维生素C / E", f"{m['vc']} / {m['ve']} mg"),
    ]
    for k, v in rows:
        c = t.add_row().cells
        c[0].text, c[1].text = k, v

    doc.add_heading("建议", level=2)
    for s in [
        "进餐顺序：先凉拌菜打底 → 蛋类 → 窝窝头 → 最后喝豆浆，蛋白先行有利于平稳餐后血糖。",
        "减脂控糖：整餐碳水35g、供能比38%，对血糖友好；凉拌菜少淋油的话脂肪还能再降3-4g。",
        "优化方向：左下格可补一份绿叶菜，膳食纤维从5g提到8g左右；蛋类吃法保持1蛋黄+2蛋白即可。",
    ]:
        doc.add_paragraph(s, style="List Bullet")

    doc.save(REC_DOCX)
    print("REC DOCX OK:", REC_DOCX)


if __name__ == "__main__":
    build_db()
    build_records_xlsx()
    build_records_docx()
    print("ALL DONE")
